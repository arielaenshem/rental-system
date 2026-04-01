#!/usr/bin/env python3
"""
Generate a sample customers.xlsx template file with Hebrew headers and example rows.
Run once:  python create_template.py
"""
import os
from datetime import date, timedelta

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_template(output_path: str = "data/customers.xlsx") -> None:
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    today = date.today()
    example_rows = [
        {
            "customer_name":       "ישראל ישראלי",
            "expected_amount":     5000.00,
            "ownership_pct":       100.0,
            "cpi_base_date":       date(2023, 1, 1),
            "contract_start_date": date(2023, 1, 1),
            "contract_end_date":   date(2026, 12, 31),
            "bank_account":        "123456",
            "customer_email":      "israel@example.com",
            "notes":               "דירה ברחוב הרצל 10",
        },
        {
            "customer_name":       "רחל כהן",
            "expected_amount":     3500.00,
            "ownership_pct":       50.0,
            "cpi_base_date":       date(2022, 6, 1),
            "contract_start_date": date(2022, 6, 1),
            "contract_end_date":   today + timedelta(days=20),  # triggers renewal alert
            "bank_account":        "654321,111222",
            "customer_email":      "rachel@example.com",
            "notes":               "בעלות משותפת – 50%",
        },
    ]

    # Hebrew display headers (also accepted as aliases by the parser)
    col_map = {
        "customer_name":       "שם לקוח",
        "expected_amount":     "סכום צפוי",
        "ownership_pct":       "אחוז החזקה",
        "cpi_base_date":       "תאריך בסיס מדד",
        "contract_start_date": "תאריך התחלה",
        "contract_end_date":   "תאריך סיום",
        "bank_account":        "חשבון בנק",
        "customer_email":      "מייל לקוח",
        "notes":               "הערות",
    }

    df = pd.DataFrame(example_rows).rename(columns=col_map)
    df.to_excel(output_path, index=False, sheet_name="לקוחות")

    # Style the workbook
    wb = load_workbook(output_path)
    ws = wb.active

    header_fill  = PatternFill("solid", fgColor="1F4E79")
    header_font  = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    border_side  = Side(style="thin", color="BFBFBF")
    cell_border  = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align  = Alignment(horizontal="right",  vertical="center", wrap_text=True)

    for col_idx, cell in enumerate(ws[1], start=1):
        cell.fill      = header_fill
        cell.font      = header_font
        cell.alignment = center_align
        cell.border    = cell_border
        ws.column_dimensions[get_column_letter(col_idx)].width = 20

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border    = cell_border
            cell.alignment = right_align

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # Add a second sheet with instructions
    ws_info = wb.create_sheet("הוראות")
    instructions = [
        ["עמודה",          "תיאור",                                                 "דוגמה"],
        ["שם לקוח",        "שם מלא של הלקוח",                                     "ישראל ישראלי"],
        ["סכום צפוי",      "הסכום החודשי הצפוי (סטייה של 5% מותרת)",             "5000"],
        ["אחוז החזקה",     "אחוז הבעלות לחישוב החשבונית (0-100)",                 "100"],
        ["תאריך בסיס מדד", "תאריך בסיס להצמדה למדד המחירים לצרכן (DD/MM/YYYY)", "01/01/2023"],
        ["תאריך התחלה",    "תאריך תחילת החוזה",                                   "01/01/2023"],
        ["תאריך סיום",     "תאריך סיום החוזה – תישלח התראה 30 יום לפני",         "31/12/2026"],
        ["חשבון בנק",      "מספר חשבון בנק דיסקונט (ניתן להזין כמה מופרדים בפסיק)", "123456"],
        ["מייל לקוח",      "כתובת מייל של הלקוח (אופציונלי)",                    "a@b.com"],
        ["הערות",          "הערות חופשיות שיופיעו בחשבונית (אופציונלי)",          "דירה ברחוב הרצל"],
    ]
    for row in instructions:
        ws_info.append(row)
    for cell in ws_info[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="2E75B6")
        cell.font = Font(bold=True, color="FFFFFF")
    for col in ["A", "B", "C"]:
        ws_info.column_dimensions[col].width = 35

    wb.save(output_path)
    print(f"Template created: {output_path}")


if __name__ == "__main__":
    create_template()
